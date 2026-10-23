"""右侧预览窗格: 图片 文本 表格, 未知类型尝试以文本模式预览

外观与设计稿一致: 276 像素宽的一列, 左边缘一条 1 像素分隔线, 自上而下是
小节标题 预览框 属性字段与底部的两个按钮. 可预览的类型见 constants, 与
Windows 资源管理器的预览窗格覆盖范围一致
"""

import csv
import os
import tkinter as tk
from tkinter import ttk

from PIL import ImageTk

from .. import kit
from ..dpi import scale
from ..errors import open_path
from ..imaging import load_thumbnail, make_gradient
from ..theme import SIZE_BODY, SIZE_LABEL, SIZE_SECTION, mono, tint
from .constants import (
    IMAGE_EXTS,
    MAX_PREVIEW_BYTES,
    MAX_PREVIEW_ROWS,
    PREVIEW_WIDTH,
    TABLE_EXTS,
    TEXT_EXTS,
)

#: 预览框的高度, 与设计稿一致
PREVIEW_BOX_HEIGHT = scale(196)

#: 属性字段的默认内容
DEFAULT_FIELDS = (("name", "文件名"), ("size", "大小"))


class PreviewPane(tk.Frame):
    """文件预览窗格; 通过 show(path) 展示内容"""

    def __init__(self, master, theme, fields=DEFAULT_FIELDS, on_open=None,
                 on_reveal=None, width=PREVIEW_WIDTH):
        super().__init__(master, bd=0, highlightthickness=0, width=width)
        self.theme = theme
        tint(self, bg="panel")
        self.configure(bg=theme.rgb("panel"))
        self.pack_propagate(False)
        self._preview_path = None
        self._preview_img = None
        self._gradient = None
        self._on_open = on_open or self._open_default
        self._on_reveal = on_reveal or self._reveal_default
        self._fields = tuple(fields)
        self._info_provider = None

        sep = kit.hairline(self, theme, orient="vertical")
        sep.pack(side="left", fill="y")

        inner = tk.Frame(self, bd=0, highlightthickness=0)
        tint(inner, bg="panel")
        inner.configure(bg=theme.rgb("panel"))
        inner.pack(side="left", fill="both", expand=True, padx=scale(14),
                   pady=scale(14))

        kit.text(inner, self.theme, value="预览", size=SIZE_SECTION,
                 token="muted", bold=True,
                 surface="panel").pack(anchor="w", pady=(0, scale(10)))

        self.box = tk.Frame(inner, bd=0, highlightthickness=1,
                            height=PREVIEW_BOX_HEIGHT)
        tint(self.box, bg="fill", hl="hair", hlc="hair")
        self.box.configure(bg=theme.rgb("fill"),
                           highlightbackground=theme.rgb("hair"),
                           highlightcolor=theme.rgb("hair"))
        self.box.pack(fill="x")
        self.box.pack_propagate(False)

        holder = tk.Frame(inner, bd=0, highlightthickness=0)
        tint(holder, bg="panel")
        holder.configure(bg=theme.rgb("panel"))
        holder.pack(fill="x", pady=(scale(12), 0))
        self._grid, self._labels, self._apply_values = kit.value_grid(
            holder, theme, self._fields, surface="panel")
        self._grid.pack(fill="x")

        actions = tk.Frame(inner, bd=0, highlightthickness=0)
        tint(actions, bg="panel")
        actions.configure(bg=theme.rgb("panel"))
        actions.pack(side="bottom", anchor="w")
        self.btn_open = ttk.Button(actions, text="打开", style="TButton",
                                   command=self._open_preview_file)
        self.btn_open.pack(side="left")
        self.btn_reveal = ttk.Button(actions, text="在文件夹中显示",
                                     style="TButton",
                                     command=self._reveal_preview_file)
        self.btn_reveal.pack(side="left", padx=(scale(8), 0))

        self.show(None)

    # ---------------- 对外接口 ----------------
    @property
    def path(self):
        """当前预览的文件路径"""
        return self._preview_path

    def show(self, path):
        """根据文件路径刷新预览; path 为空时显示占位内容

        @param path: 文件路径
        """
        self._preview_path = path
        self._preview_img = None
        self._clear_box()

        if not path or not os.path.isfile(path):
            self._apply_values({key: "-" for key, _caption in self._fields})
            self._placeholder_gradient()
            return

        info = {"name": _shorten(os.path.basename(path)),
                "size": _format_size(self._file_size(path))}
        info.update(self.extra_fields(path))
        if self._info_provider is not None:
            info.update(self._info_provider(path) or {})
        self._apply_values(info)

        ext = os.path.splitext(path)[1].lower()
        if ext in IMAGE_EXTS:
            self._preview_image(path)
        elif ext in TEXT_EXTS:
            self._preview_text(path)
        elif ext in TABLE_EXTS:
            self._preview_table(path, ext)
        else:
            message = f"暂不支持预览 {ext} 文件" if ext else "暂不支持预览该文件"
            self._try_text_fallback(path, message)
        if not self.box.winfo_children():
            self._placeholder_gradient()

    def extra_fields(self, path):
        """子类或页面补充的属性字段

        @param path: 文件路径
        @return: {字段键: 取值}
        """
        return {}

    def set_info_provider(self, provider):
        """登记属性字段的取值来源

        页面可以借此补充 尺寸 与 相似度 这类需要额外查询的字段

        @param provider: callable(path) -> {字段键: 取值}
        """
        self._info_provider = provider
        self.show(self._preview_path)

    def refresh(self):
        """使用当前主题重新绘制正在显示的预览内容"""
        self._gradient = None
        self.show(self._preview_path)

    # ---------------- 预览框内容 ----------------
    def _clear_box(self):
        for child in self.box.winfo_children():
            child.destroy()

    def _placeholder_gradient(self):
        """没有可预览的内容时用设计稿的渐变填满预览框"""
        if self._gradient is None:
            palette = self.theme.palette
            image = make_gradient(
                (PREVIEW_WIDTH - 32, PREVIEW_BOX_HEIGHT - 2),
                _rgb(palette["ph1"]), _rgb(palette["ph2"]))
            self._gradient = ImageTk.PhotoImage(image)
        label = tk.Label(self.box, image=self._gradient, bd=0)
        tint(label, bg="fill")
        label.configure(bg=self.theme.rgb("fill"))
        label.pack(fill="both", expand=True)

    def _placeholder_text(self, message):
        label = kit.text(self.box, self.theme, value=message, size=SIZE_BODY,
                         token="muted", surface="fill", anchor="center",
                         wrap=PREVIEW_WIDTH - 60)
        label.pack(fill="both", expand=True)

    def _preview_image(self, path):
        self.update_idletasks()
        width = max(self.box.winfo_width() - 20, 80)
        height = max(PREVIEW_BOX_HEIGHT - 20, 60)
        img = load_thumbnail(path, (width, height))
        if img is None:
            self._try_text_fallback(path, "图片加载失败")
            return
        self._preview_img = ImageTk.PhotoImage(img)
        label = tk.Label(self.box, image=self._preview_img, bd=0)
        tint(label, bg="fill")
        label.configure(bg=self.theme.rgb("fill"))
        label.pack(fill="both", expand=True)

    def _open_preview_file(self):
        if self._preview_path:
            self._on_open(self._preview_path)

    def _reveal_preview_file(self):
        if self._preview_path:
            self._on_reveal(self._preview_path)

    @staticmethod
    def _open_default(path):
        open_path(path)

    @staticmethod
    def _reveal_default(path):
        open_path(path, reveal=True)

    def _preview_text(self, path):
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                content = handle.read(MAX_PREVIEW_BYTES)
        except Exception:
            self._try_text_fallback(path, "无法读取文本文件")
            return
        self._show_text_content(content)

    def _show_text_content(self, content):
        text = tk.Text(self.box, wrap="none", bd=0, highlightthickness=0,
                       font=mono(SIZE_LABEL), padx=scale(10),
                       pady=scale(8))
        tint(text, bg="fill", fg="text", insert="text", sel_bg="accent",
             sel_fg="on_accent")
        text.configure(bg=self.theme.rgb("fill"), fg=self.theme.rgb("text"),
                       selectbackground=self.theme.rgb("accent"),
                       selectforeground=self.theme.rgb("on_accent"))
        text.insert("1.0", content)
        if len(content) >= MAX_PREVIEW_BYTES:
            text.insert("end", "\n\n... 文件较大, 仅预览前 200KB")
        text.configure(state="disabled")
        text.pack(fill="both", expand=True)

    def _try_text_fallback(self, path, fail_message):
        """无法按原类型预览时, 尝试把文件当作文本读取

        @param path: 文件路径
        @param fail_message: 确实读不出来时显示的说明
        """
        try:
            with open(path, "rb") as handle:
                data = handle.read(MAX_PREVIEW_BYTES)
        except Exception:
            self._placeholder_text(fail_message)
            return
        if _looks_binary(data):
            self._placeholder_text(fail_message)
            return
        self._show_text_content(data.decode("utf-8", errors="replace"))

    def _preview_table(self, path, ext):
        if ext == ".xlsx":
            rows = self._read_xlsx(path)
        else:
            rows = []
            try:
                with open(path, "r", encoding="utf-8-sig", errors="replace",
                          newline="") as handle:
                    dialect = "excel-tab" if ext == ".tsv" else "excel"
                    for row in csv.reader(handle, dialect):
                        rows.append(row)
                        if len(rows) >= MAX_PREVIEW_ROWS:
                            break
            except Exception:
                self._try_text_fallback(path, "无法读取表格文件")
                return
        if rows is None:
            return  # 读取失败时已显示提示
        content = "\n".join("\t".join(cell for cell in row) for row in rows)
        self._show_text_content(content)

    def _read_xlsx(self, path):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self._placeholder_text("未安装 openpyxl, 无法预览 Excel 表格")
            return None
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book[book.sheetnames[0]]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([_cell_text(value) for value in row])
                if len(rows) >= MAX_PREVIEW_ROWS:
                    break
            book.close()
            return rows
        except Exception:
            self._try_text_fallback(path, "Excel 文件读取失败")
            return None

    # ---------------- 工具 ----------------
    @staticmethod
    def _file_size(path):
        try:
            return os.path.getsize(path)
        except OSError:
            return 0


def _rgb(text):
    """把 #RRGGBB 换成 (r, g, b) 三元组

    @param text: 颜色字符串
    @return: 三元组
    """
    return (int(text[1:3], 16), int(text[3:5], 16), int(text[5:7], 16))


def _looks_binary(data):
    """粗略判断一段字节是不是二进制内容

    @param data: 文件开头的一段字节
    @return: True 表示不适合当作文本预览
    """
    if not data:
        return False
    sample = data[:8000]
    if b"\x00" in sample:
        return True
    control = sum(1 for byte in sample
                  if byte < 32 and byte not in (9, 10, 13, 12, 27))
    return control / len(sample) > 0.30


def _cell_text(value):
    """表格单元格转成一行文本

    @param value: 单元格取值
    @return: 展示用文本
    """
    if value is None:
        return ""
    return str(value).replace("\n", " ")[:60]


def _format_size(size):
    """把字节数格式化成便于阅读的大小

    @param size: 字节数
    @return: 展示用文本
    """
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / 1024 / 1024:.1f} MB"


def _shorten(name, limit=32):
    """标题一行放不下时保留尾部, 便于看出真实文件名

    @param name: 文件名
    @param limit: 最大字符数
    @return: 截断后的文件名
    """
    return name if len(name) <= limit else "..." + name[-(limit - 3):]
